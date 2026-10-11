import json,sys,shutil
from pptx import Presentation
a=json.load(sys.stdin)
shutil.copyfile(a["template"],a["output"])
prs=Presentation(a["output"])
prs.slides[0].shapes.title.text=a["cover"]
s=prs.slides[a["slide"]]
s.shapes.title.text=a["tableTitle"]
old=next(sh for sh in s.shapes if sh.has_table)
left,top,width,height=old.left,old.top,old.width,old.height
old._element.getparent().remove(old._element)
rows=[["제품","수량","매출(세전)"]]+[[r["item"],str(r["qty"]),"{:,}".format(r["ex"])] for r in a["rows"]]
newtbl=s.shapes.add_table(len(rows),len(rows[0]),left,top,width,height).table
for i,r in enumerate(rows):
    for j,v in enumerate(r): newtbl.cell(i,j).text=v
prs.save(a["output"])
total=sum(r["ex"] for r in a["rows"])
print(json.dumps({"summary":"render %s total %d"%(a["cover"],total),"json":{"total":total,"title":a["cover"]}}))